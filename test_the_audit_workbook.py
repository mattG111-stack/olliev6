"""The live data as a workbook, with the checks already run.

    "also have an exel download for all live data so i can download it and you
     audit it for bugs"

The value is not the download, it is the first sheet. Every fault this system
has shipped was visible in the data at the time — 1,436 of 1,586 listings
"advertising" a price in a market where four in five name none; a headline deal
of +2,296% against a council record that had valued the land and not the house;
672 listings live where 9,281 had been the week before. Each was there to be
read and nobody could read it, because reading it meant knowing which column to
look down and what the number should have been.

So these tests do not check that a file appears. They rebuild those three faults
and assert the sheet SAYS SO — and, just as important, that a healthy batch is
not covered in warnings, because a check that cries wolf is a check nobody
finishes reading.
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import pandas as pd
import pytest

from app import audit_export
from app.ingest import ingest_for_sale, ingest_sold
from app.models import BatchType, ImportBatch, PropertyForSale
from app.release import publish_release


def _sold() -> pd.DataFrame:
    return pd.DataFrame([{
        "address": f"{i} Sold Street", "suburb": "Papakura",
        "district": "Papakura", "region": "Auckland", "property_type": "House",
        "key_bedrooms": 3, "key_bathrooms": 1,
        "key_floor_area": f"{130 + i} sqm", "key_land_area": f"{700 + i * 10} sqm",
        "cv_numeric": 1_000_000, "price_numeric": 1_050_000 + i * 10_000,
        "sale_price": 1_050_000 + i * 10_000,
        "land_value_numeric": 700_000, "improvement_value_numeric": 300_000,
        "type_of_title": "Freehold", "sold_date": "2026-08-01",
    } for i in range(12)])


def _for_sale(slugs, **over) -> pd.DataFrame:
    base = dict(district="Papakura", region="Auckland", property_type="House",
                suburb="Papakura", price_display="$1,100,000",
                price_numeric=1_100_000, sale_method="fixed price",
                image_1_url="https://cdn.example/x.jpg",
                cv_numeric=1_000_000, land_value_numeric=700_000,
                improvement_value_numeric=300_000, zoning="Mixed Housing Urban",
                key_bedrooms=3, key_bathrooms=1, key_floor_area=140,
                key_land_area=750, type_of_title="Freehold")
    base.update(over)
    return pd.DataFrame([{**base, "address": f"{s} Example Road", "slug_id": s,
                          "url": f"https://a-portal.example/listing/{s}"}
                         for s in slugs])


@pytest.fixture()
def live(db_session):
    ingest_sold(db_session, _sold(), "sold.csv", region="Auckland", publish=False)
    ingest_for_sale(db_session, _for_sale([f"h-{i}" for i in range(20)],
                                          key_land_area=1600), _sold(),
                    "week1.csv", region="Auckland", publish=False,
                    fill_missing=False)
    publish_release(db_session, region="Auckland")
    return db_session


def _checks(db) -> dict[str, tuple[str, str]]:
    """{question: (answer, verdict)} — read back out of the real workbook."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(audit_export.build(db, "Auckland")))
    ws = wb["Checks"]
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row and row[0]:
            out[str(row[0])] = (str(row[1]), str(row[3] or ""))
    return out


def _flagged(db) -> set[str]:
    return {q for q, (_, v) in _checks(db).items() if v == "LOOK AT THIS"}


def _live_rows(db):
    b = (db.query(ImportBatch)
         .filter(ImportBatch.batch_type == BatchType.FOR_SALE.value,
                 ImportBatch.is_active.is_(True))
         .order_by(ImportBatch.id.desc()).first())
    return (db.query(PropertyForSale)
            .filter(PropertyForSale.import_batch_id == b.id).all())


# ---- it is a real workbook ---------------------------------------------------
def test_it_opens_in_excel(live):
    """openpyxl writing something is not the same as Excel reading it."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(audit_export.build(live, "Auckland")))
    assert wb.sheetnames == ["Checks", "Listings"]
    assert wb["Listings"].max_row == 21          # 20 listings and a heading


def test_every_listing_is_in_it(live):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(audit_export.build(live, "Auckland")))
    ws = wb["Listings"]
    heads = [c.value for c in ws[1]]
    refs = {r[heads.index("Listing ref")] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert refs == {f"h-{i}" for i in range(20)}


def test_an_empty_account_does_not_crash(db_session):
    """The first thing a new operator would press."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(audit_export.build(db_session, "Auckland")))
    assert wb["Listings"].max_row == 1


def test_dates_are_readable(live):
    """Day, month, year — never the ISO string."""
    from openpyxl import load_workbook

    for p in _live_rows(live):
        p.link_checked_at = datetime(2026, 3, 9, tzinfo=timezone.utc)
    live.commit()

    wb = load_workbook(io.BytesIO(audit_export.build(live, "Auckland")))
    ws = wb["Listings"]
    heads = [c.value for c in ws[1]]
    col = heads.index("Link last checked")
    assert ws.cell(row=2, column=col + 1).value == "09/03/2026"


# ---- a healthy batch is quiet -----------------------------------------------
def test_only_the_things_actually_wrong_are_flagged(live):
    """A check that cries wolf is a check nobody finishes reading.

    Three flags are expected on this fixture and every one is CORRECT — which
    is the point, so they are named rather than waved through:

      every listing carries a price, which is the fault that shipped
      no link has been checked, because the daily sweep has not run
      nothing is subdividable, because twenty synthetic Papakura sections give
      the engine no section sales to rate land from

    Anything beyond these three is the sheet inventing a problem.
    """
    assert _flagged(live) == {
        "Vendor actually named a price",
        "Link never checked",
        "Can be subdivided",
    }


# ---- and each fault we have actually shipped is called out -------------------
def test_it_catches_a_price_on_every_listing(live):
    """THE ONE THAT SHIPPED. 1,436 of 1,586 listings 'asking' a price, when the
    true figure was 308. Anybody who knows this market stops at 91%; nobody was
    shown the number."""
    flagged = _flagged(live)
    assert "Vendor actually named a price" in flagged
    assert "100%" in _checks(live)["Vendor actually named a price"][0]


def test_a_normal_market_does_not_trip_that_check(db_session):
    """Four in five selling by auction or negotiation is the NORMAL case and
    must read as normal, or the check is noise."""
    ingest_sold(db_session, _sold(), "sold.csv", region="Auckland", publish=False)
    df = pd.concat([
        _for_sale([f"p-{i}" for i in range(4)]),
        _for_sale([f"a-{i}" for i in range(16)], sale_method="auction",
                  price_display="Auction", price_numeric=None),
    ], ignore_index=True)
    ingest_for_sale(db_session, df, _sold(), "week1.csv", region="Auckland",
                    publish=False, fill_missing=False)
    publish_release(db_session, region="Auckland")
    assert "Vendor actually named a price" not in _flagged(db_session)


def test_it_catches_a_value_that_ran_away(live):
    """+2,296% against a council record that valued the land and not the house.
    The deal page picks its hero by the largest margin it can find, so the most
    broken row in the batch becomes the headline."""
    p = _live_rows(live)[0]
    p.fair_value = (p.asking_price or 1_100_000) * 3
    live.commit()
    assert "Our value more than double the asking" in _flagged(live)


def test_it_catches_the_same_house_twice(live):
    """One house on the site twice, at two different prices."""
    p = _live_rows(live)[0]
    live.add(PropertyForSale(import_batch_id=p.import_batch_id, slug_id=p.slug_id,
                             address=p.address, suburb=p.suburb,
                             floor_area_m2=140.0,
                             is_underpriced=False, is_cashflow_positive=False,
                             is_subdividable=False))
    live.commit()
    assert "The same listing more than once" in _flagged(live)


def test_it_catches_a_book_the_link_check_never_reaches(live):
    """The fault found this week: the sweep restarted at the top every upload
    and never got to the bottom. It ran, reported hundreds of checks, and was
    re-checking the same front slice for ever."""
    assert "Link never checked" in _flagged(live)


def test_it_catches_listings_with_no_link(live):
    """The link check cannot open what it has no address for, so these can never
    leave the site — the book would grow for ever."""
    for p in _live_rows(live)[:3]:
        p.url = None
    live.commit()
    assert "No link to the advertisement" in _flagged(live)


# ---- the hidden rows are where faults live ----------------------------------
def test_a_hidden_row_says_why_it_is_hidden(live):
    """A flag says something is wrong somewhere. A reason says what."""
    from openpyxl import load_workbook

    p = _live_rows(live)[0]
    p.is_held = True
    p.hold_reason = "Below $10,000 margin"
    live.commit()

    wb = load_workbook(io.BytesIO(audit_export.build(live, "Auckland")))
    ws = wb["Listings"]
    heads = [c.value for c in ws[1]]
    rows = {r[heads.index("Listing ref")]: r for r in
            ws.iter_rows(min_row=2, values_only=True)}
    row = rows[p.slug_id]
    assert row[heads.index("Visible on site")] == "NO"
    assert "Below $10,000 margin" in row[heads.index("Why not visible")]


def test_a_row_a_customer_cannot_see_is_still_in_the_file(live):
    """The whole reason for including them: a held row is invisible on the site,
    so a fault that only affects held rows is invisible everywhere."""
    from openpyxl import load_workbook

    for p in _live_rows(live):
        p.is_held = True
    live.commit()

    wb = load_workbook(io.BytesIO(audit_export.build(live, "Auckland")))
    assert wb["Listings"].max_row == 21, "the hidden rows were dropped"
    assert _checks(live)["Listings a customer can see"][0] == "0"


def test_a_dead_link_is_named_as_the_reason(live):
    p = _live_rows(live)[0]
    p.link_dead_at = datetime.now(timezone.utc) - timedelta(days=2)
    live.commit()
    why = audit_export._visible_reason(p)
    assert why and "Advertisement is gone" in why
